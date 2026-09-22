#!/usr/bin/env python3
"""Export an html-slides deck (index.html) to pixel-perfect PDF and editable PPTX.

PDF  : per-page Playwright screenshots (2x) assembled with img2pdf.
PPTX : per-page background screenshot with text made invisible
       (-webkit-text-fill-color: transparent, keeps currentColor decorations),
       plus editable textboxes rebuilt from DOM geometry (python-pptx).

Usage:
  python export_deck.py <deck_dir> [--pdf] [--pptx] [--name 输出名]
  (neither flag = export both; requires conda env with playwright/pptx/img2pdf)
"""
import argparse
import http.server
import socketserver
import threading
from pathlib import Path

STAGE_W, STAGE_H = 1280, 720
EMU_PER_PX = 9525          # 12192000 EMU / 1280 px  (== 96 dpi)
PT_PER_PX = 0.75

FREEZE_CSS = """
*,*::before,*::after{animation:none!important;transition:none!important}
[data-a]{opacity:1!important;transform:none!important}
.ln{stroke-dashoffset:0!important}
#nav,#scrub,#scrubGlow,#tip,#track,#ticks{display:none!important}
"""

EXTRACT_JS = r"""
() => {
  const stage = document.querySelector('#stage') || document.body;
  const srect = stage.getBoundingClientRect();
  const slide = document.querySelector('.slide.active');
  const blocks = [];
  const norm = t => t.replace(/\s+/g, ' ');
  // inline 与否按 computed display 判断：flex 子项 / absolute 定位的 span 会被块化，须各自成框
  const isInline = el => el.tagName !== 'svg' && el.tagName !== 'SVG' &&
                         getComputedStyle(el).display.startsWith('inline');
  const alphaOf = c => {
    const m = (c || '').match(/rgba?\(([^)]+)\)/);
    if (!m) return 1;
    const p = m[1].split(',');
    return p.length > 3 ? parseFloat(p[3]) : 1;
  };
  function pseudoRun(el, which) {                 // ::before/::after 的文字（如 kicker 的 " /"）
    const c = getComputedStyle(el, which);
    let t = c.content;
    if (!t || t === 'none' || t === 'normal') return null;
    if (t.startsWith('"') || t.startsWith("'")) t = t.slice(1, -1);
    if (!t.trim() || alphaOf(c.color) === 0) return null;
    return {text: t, bold: parseInt(c.fontWeight) >= 600, italic: c.fontStyle === 'italic',
            color: c.color, size: parseFloat(c.fontSize), family: c.fontFamily,
            ls: parseFloat(c.letterSpacing) || 0};
  }
  function runsOf(el) {
    const runs = [];
    const push = (text, node) => {
      if (!text) return;
      const c = getComputedStyle(node.nodeType === 1 ? node : el);
      // color 全透明的文字（-webkit-text-stroke 描边字等）留在底图，不生成文本框
      if (alphaOf(c.color) === 0) return;
      if (node.nodeType === 1) { const pb = pseudoRun(node, '::before'); if (pb) runs.push(pb); }
      runs.push({text, bold: parseInt(c.fontWeight) >= 600,
                 italic: c.fontStyle === 'italic', color: c.color,
                 size: parseFloat(c.fontSize), family: c.fontFamily,
                 ls: parseFloat(c.letterSpacing) || 0});
      if (node.nodeType === 1) { const pa = pseudoRun(node, '::after'); if (pa) runs.push(pa); }
    };
    const pb = pseudoRun(el, '::before'); if (pb) runs.push(pb);
    for (const n of el.childNodes) {
      if (n.nodeType === 3) push(norm(n.textContent), n);
      else if (n.nodeType === 1 && n.tagName === 'BR') runs.push({text: '\n', br: true});
      else if (n.nodeType === 1 && isInline(n)) push(norm(n.textContent), n);
    }
    const pa = pseudoRun(el, '::after'); if (pa) runs.push(pa);
    while (runs.length && !runs[0].br && !runs[0].text.trim()) runs.shift();
    while (runs.length && !runs.at(-1).br && !runs.at(-1).text.trim()) runs.pop();
    return runs;
  }
  function inlineRect(el) {                       // union rect of own text + inline children
    const r = document.createRange();
    let rect = null;
    for (const n of el.childNodes) {
      if ((n.nodeType === 3 && n.textContent.trim()) ||
          (n.nodeType === 1 && isInline(n) && n.textContent.trim())) {
        r.selectNodeContents(n);
        const b = r.getBoundingClientRect();
        if (b.width < 1 || b.height < 1) continue;
        rect = rect ? {left: Math.min(rect.left, b.left), top: Math.min(rect.top, b.top),
                       right: Math.max(rect.right, b.right), bottom: Math.max(rect.bottom, b.bottom)}
                    : {left: b.left, top: b.top, right: b.right, bottom: b.bottom};
      }
    }
    return rect;
  }
  function emit(el, rect) {
    const cs = getComputedStyle(el);
    const runs = runsOf(el);
    if (!runs.length || !runs.some(r => r.text.trim())) return;
    blocks.push({x: rect.left - srect.left, y: rect.top - srect.top,
                 w: rect.right - rect.left, h: rect.bottom - rect.top,
                 align: cs.textAlign, size: parseFloat(cs.fontSize),
                 lh: parseFloat(cs.lineHeight) || parseFloat(cs.fontSize) * 1.4,
                 family: cs.fontFamily, nowrap: cs.whiteSpace === 'nowrap', runs});
    el.setAttribute('data-xtext', '1');
  }
  function walk(el) {
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden' || parseFloat(cs.opacity) === 0) return;
    const own = [...el.childNodes].some(n =>
      (n.nodeType === 3 && n.textContent.trim()) ||
      (n.nodeType === 1 && isInline(n) && n.textContent.trim()));
    const blockKids = [...el.children].filter(c => !isInline(c) && c.tagName !== 'BR');
    if (own) {
      if (!blockKids.length) {
        // 无块级子元素：用 content box 定位，padding 里的装饰（li 菱形、quote 竖线）才不会被文字压住
        const r = el.getBoundingClientRect();
        const rect = {left: r.left + parseFloat(cs.paddingLeft) + parseFloat(cs.borderLeftWidth),
                      top: r.top + parseFloat(cs.paddingTop) + parseFloat(cs.borderTopWidth),
                      right: r.right - parseFloat(cs.paddingRight) - parseFloat(cs.borderRightWidth),
                      bottom: r.bottom - parseFloat(cs.paddingBottom) - parseFloat(cs.borderBottomWidth)};
        // 非居中块的起点用文字实际渲染位置（Range 实测）——flex 居中（如数字方块）不体现在
        // text-align 里，content box 起点会让文字漂移；居中块保留全宽让 PPT 重新居中
        if (!cs.textAlign.includes('center')) {
          const ir = inlineRect(el);
          if (ir) { rect.left = ir.left; rect.top = Math.max(rect.top, ir.top); }
        }
        if (rect.right - rect.left > 1 && rect.bottom - rect.top > 1) emit(el, rect);
      } else {
        const r = inlineRect(el);                 // mixed: own inline text + block children
        if (r) emit(el, r);
      }
    }
    for (const c of blockKids) if (c.tagName !== 'SVG' && c.tagName !== 'svg') walk(c);
  }
  walk(slide);
  const chrome = document.querySelector('#chrome');
  if (chrome) walk(chrome);
  return blocks;
}
"""

PREP_JS = r"""
async () => {
  document.querySelectorAll('[data-src]').forEach(el => {
    el.src = el.dataset.src; el.removeAttribute('data-src');
    if (el.tagName === 'VIDEO') el.preload = 'auto';
  });
  document.querySelectorAll('video').forEach(v => v.removeAttribute('controls'));
  await document.fonts.ready;
  const media = [...document.querySelectorAll('.slide.active img, .slide.active video')];
  await Promise.all(media.map(m => new Promise(res => {
    const done = () => res();
    if (m.tagName === 'IMG') { if (m.complete) done(); else { m.onload = m.onerror = done; } }
    else { if (m.readyState >= 2) done(); else { m.onloadeddata = m.onerror = done; setTimeout(done, 4000); } }
  })));
}
"""


def serve(root: Path):
    handler = type('H', (http.server.SimpleHTTPRequestHandler,),
                   {'directory': str(root), 'log_message': lambda *a: None})
    handler = lambda *a, **kw: http.server.SimpleHTTPRequestHandler(*a, directory=str(root), **kw)
    srv = socketserver.TCPServer(('127.0.0.1', 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, srv.server_address[1]


def font_name(family: str) -> str:
    f = (family or '').lower()
    if any(k in f for k in ('mono', 'menlo', 'courier', 'consolas')):
        return 'Menlo'
    if any(k in f for k in ('songti', 'stsong', 'serif', 'song')):
        return 'Songti SC'
    return 'PingFang SC'


def build_pptx(shots, blocks_all, out: Path):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    ALIGN = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT,
             'justify': PP_ALIGN.JUSTIFY, 'left': PP_ALIGN.LEFT, 'start': PP_ALIGN.LEFT}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(STAGE_W * EMU_PER_PX), Emu(STAGE_H * EMU_PER_PX)
    blank = prs.slide_layouts[6]

    for png, blocks in zip(shots, blocks_all):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        for b in blocks:
            pad = 4
            bx, bw = b['x'], b['w']
            if b.get('nowrap'):                     # 双保险：wrap="none" 之外再加宽（LibreOffice 会忽略 wrap 属性）
                extra = bw * 0.3 + 12
                if 'center' in (b.get('align') or ''):
                    bx -= extra / 2
                bw += extra
            tb = slide.shapes.add_textbox(
                Emu(max(0, int((bx - pad) * EMU_PER_PX))),
                Emu(max(0, int((b['y'] - pad / 2) * EMU_PER_PX))),
                Emu(int((bw + 2 * pad) * EMU_PER_PX)),
                Emu(int((b['h'] + pad) * EMU_PER_PX)))
            tf = tb.text_frame
            # HTML 里 nowrap 的块（流程图节点等）不许 PPT 换行——PPT 字形略宽，换行会把末字挤出框
            tf.word_wrap = not b.get('nowrap')
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            para = tf.paragraphs[0]
            para.alignment = ALIGN.get(b['align'], PP_ALIGN.LEFT)
            if b.get('lh') and b.get('size'):
                para.line_spacing = Pt(b['lh'] * PT_PER_PX)
            for r in b['runs']:
                if r.get('br'):
                    para = tf.add_paragraph()
                    para.alignment = ALIGN.get(b['align'], PP_ALIGN.LEFT)
                    if b.get('lh'):
                        para.line_spacing = Pt(b['lh'] * PT_PER_PX)
                    continue
                run = para.add_run()
                run.text = r['text']
                run.font.size = Pt(round(r['size'] * PT_PER_PX * 2) / 2)
                run.font.bold = bool(r['bold'])
                run.font.italic = bool(r.get('italic'))
                run.font.name = font_name(r.get('family') or b.get('family'))
                c = r['color']
                if c.startswith('rgb'):
                    parts = [int(float(v)) for v in c[c.index('(') + 1:c.index(')')].split(',')[:3]]
                    run.font.color.rgb = RGBColor(*parts)
                ls = r.get('ls') or 0                       # letter-spacing → OOXML spc（1/100 pt）
                if ls:
                    run._r.get_or_add_rPr().set('spc', str(int(round(ls * PT_PER_PX * 100))))
    prs.save(str(out))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('deck_dir')
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--pptx', action='store_true')
    ap.add_argument('--name', default=None, help='输出文件名（不含扩展名），默认 deck 目录名')
    args = ap.parse_args()
    do_pdf = args.pdf or not (args.pdf or args.pptx)
    do_pptx = args.pptx or not (args.pdf or args.pptx)

    deck = Path(args.deck_dir).resolve()
    assert (deck / 'index.html').exists(), f'{deck}/index.html 不存在'
    name = args.name or deck.name
    tmp = deck / '.export_tmp'
    tmp.mkdir(exist_ok=True)

    from playwright.sync_api import sync_playwright
    srv, port = serve(deck)
    base = f'http://127.0.0.1:{port}/index.html'
    pdf_shots, bg_shots, blocks_all = [], [], []
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={'width': STAGE_W, 'height': STAGE_H},
                                    device_scale_factor=2)
            page.goto(base)
            n_slides = page.evaluate("document.querySelectorAll('.slide').length")
            print(f'共 {n_slides} 页')
            for i in range(1, n_slides + 1):
                page.goto(f'{base}#{i}')
                page.reload()
                page.add_style_tag(content=FREEZE_CSS)
                page.evaluate(PREP_JS)
                page.wait_for_timeout(120)
                clip = page.evaluate(
                    "()=>{const r=document.querySelector('#stage').getBoundingClientRect();"
                    "return {x:r.x,y:r.y,width:r.width,height:r.height}}")
                if do_pdf:
                    f = tmp / f'pdf_{i:03d}.png'
                    page.screenshot(path=str(f), clip=clip)
                    pdf_shots.append(f)
                if do_pptx:
                    blocks_all.append(page.evaluate(EXTRACT_JS))
                    page.evaluate("document.querySelectorAll('[data-xtext]').forEach(el=>"
                                  "el.style.setProperty('-webkit-text-fill-color','transparent','important'))")
                    f = tmp / f'bg_{i:03d}.png'
                    page.screenshot(path=str(f), clip=clip)
                    bg_shots.append(f)
                print(f'  第 {i}/{n_slides} 页完成', flush=True)
            browser.close()
    finally:
        srv.shutdown()

    if do_pdf:
        import img2pdf
        pdf_path = deck / f'{name}.pdf'
        layout = img2pdf.get_layout_fun((img2pdf.mm_to_pt(338.7), img2pdf.mm_to_pt(190.5)))
        with open(pdf_path, 'wb') as f:
            f.write(img2pdf.convert([str(s) for s in pdf_shots], layout_fun=layout))
        print(f'PDF  → {pdf_path} ({pdf_path.stat().st_size/1e6:.1f} MB)')
    if do_pptx:
        pptx_path = deck / f'{name}.pptx'
        build_pptx(bg_shots, blocks_all, pptx_path)
        print(f'PPTX → {pptx_path} ({pptx_path.stat().st_size/1e6:.1f} MB)')
    for f in tmp.iterdir():
        f.unlink()
    tmp.rmdir()


if __name__ == '__main__':
    main()
